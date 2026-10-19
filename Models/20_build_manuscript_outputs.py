#!/usr/bin/env python3
"""Create manuscript-ready tables and figures from harmonized benchmark outputs.

Selection rule
--------------
Track A and B finalists are selected using INTERNAL performance only. External
OAI/FracAtlas results are joined after selection and never influence ranking.
Track C zero-shot models are all reported; their uncertainty is the clustered
bootstrap CI from the single deterministic prediction run.

Track A/B values are reported as mean (SD) across five checkpoint seeds.
Domain-maintenance adaptation seeds have already been averaged within each
checkpoint seed by 18_harmonized_final_analysis.py.
"""

from __future__ import annotations

import argparse
import json
import math
import re
from pathlib import Path

import numpy as np
import pandas as pd

try:
    import matplotlib.pyplot as plt
except ImportError as exc:  # pragma: no cover
    raise SystemExit("matplotlib is required: pip install matplotlib") from exc


METRICS = [
    "accuracy",
    "balanced_accuracy_supported",
    "macro_f1_supported",
    "macro_f1_official",
    "weighted_f1_supported",
    "invalid_prediction_rate",
]

DISPLAY = {
    "tf_efficientnet_b0": "EfficientNet-B0",
    "resnet50": "ResNet-50",
    "densenet121": "DenseNet-121",
    "convnext_tiny": "ConvNeXt-Tiny",
    "vit_b16": "ViT-B/16",
    "swin_t": "Swin-T",
    "dinov2_vitb14": "DINOv2 ViT-B/14",
    "biomedclip_vitb16": "BiomedCLIP ViT-B/16",
    "clip_vitb32": "CLIP ViT-B/32",
    "qwen2_5_vl_3b": "Qwen2.5-VL-3B",
    "hgb": "Histogram gradient boosting",
    "logreg": "Logistic regression",
    "rf": "Random forest",
    "baseline": "Baseline",
    "full_ft": "Full fine-tuning",
    "linear_probe": "Linear probe",
    "maintenance": "Few-shot maintenance",
    "zero_shot": "Zero-shot",
}


def parse_args() -> argparse.Namespace:
    p = argparse.ArgumentParser(formatter_class=argparse.ArgumentDefaultsHelpFormatter)
    p.add_argument("--analysis_dir", required=True,
                   help="Directory produced by 18_harmonized_final_analysis.py")
    p.add_argument("--out_dir", required=True)
    p.add_argument("--selection_metric", default="macro_f1_supported",
                   choices=METRICS,
                   help="Internal-only metric used to select Track A/B finalists")
    p.add_argument("--primary_metric", default="macro_f1_supported", choices=METRICS)
    p.add_argument("--dpi", type=int, default=300)
    return p.parse_args()


def read_required(root: Path, name: str) -> pd.DataFrame:
    path = root / name
    if not path.exists():
        raise FileNotFoundError(f"Required harmonized output not found: {path}")
    return pd.read_csv(path, low_memory=False)


def check_columns(df: pd.DataFrame, required: list[str], label: str) -> None:
    missing = [c for c in required if c not in df.columns]
    if missing:
        raise ValueError(f"{label} is missing columns: {missing}")


def label(value) -> str:
    if pd.isna(value):
        return ""
    return DISPLAY.get(str(value), str(value))


def select_best_internal(summary: pd.DataFrame, task: str, track: str,
                         metric: str) -> pd.Series:
    split = "UPMC_internal_test" if task == "view" else "internal_test"
    candidates = summary[(summary.track == track) & (summary.task == task) &
                         (summary.split == split)].copy()
    if task == "view":
        allowed = ["baseline"] if track == "A" else ["full_ft", "linear_probe"]
    else:
        allowed = ["baseline"] if track == "A" else ["full_ft"]
    candidates = candidates[candidates.regime.isin(allowed)]
    value_col = f"{metric}_mean"
    candidates[value_col] = pd.to_numeric(candidates[value_col], errors="coerce")
    candidates = candidates.dropna(subset=[value_col])
    if candidates.empty:
        raise RuntimeError(
            f"No eligible internal candidates for Track {track} {task} on {split}"
        )
    # Explicit deterministic tie-breakers; external performance is absent here.
    order = [value_col, "model", "regime"]
    ascending = [False, True, True]
    if task == "domain":
        order.append("domain_model")
        ascending.append(True)
    return candidates.sort_values(order, ascending=ascending).iloc[0]


def selection_table(summary: pd.DataFrame, metric: str) -> pd.DataFrame:
    rows = []
    for task in ["view", "domain"]:
        for track in ["A", "B"]:
            best = select_best_internal(summary, task, track, metric)
            rows.append({
                "task": task,
                "track": track,
                "model": best.model,
                "model_display": label(best.model),
                "regime": best.regime,
                "regime_display": label(best.regime),
                "domain_model": best.get("domain_model", np.nan),
                "domain_model_display": label(best.get("domain_model", np.nan)),
                "selection_split": best.split,
                "selection_metric": metric,
                "internal_selection_value": best[f"{metric}_mean"],
                "selection_rule": "maximum internal metric; deterministic name tie-break",
            })
    return pd.DataFrame(rows)


def identity_mask(df: pd.DataFrame, selected: pd.Series) -> pd.Series:
    mask = (df.task == selected.task) & (df.track == selected.track) & (
        df.model == selected.model)
    if selected.task == "domain":
        mask &= df.domain_model.fillna("") == str(selected.domain_model or "")
    return mask


def formatted_result(row: pd.Series, metric: str, run_metrics: pd.DataFrame) -> str:
    mean = row.get(f"{metric}_mean", np.nan)
    if not np.isfinite(pd.to_numeric(mean, errors="coerce")):
        return ""
    mean = float(mean)
    if row.track in {"A", "B"}:
        sd = pd.to_numeric(row.get(f"{metric}_sd", np.nan), errors="coerce")
        return f"{mean:.3f} ({float(sd):.3f})" if np.isfinite(sd) else f"{mean:.3f}"
    match = run_metrics[(run_metrics.track == row.track) &
                        (run_metrics.task == row.task) &
                        (run_metrics.split == row.split) &
                        (run_metrics.model == row.model) &
                        (run_metrics.regime == row.regime)]
    if len(match) != 1:
        return f"{mean:.3f}"
    low = pd.to_numeric(match.iloc[0].get(f"{metric}_ci_low", np.nan), errors="coerce")
    high = pd.to_numeric(match.iloc[0].get(f"{metric}_ci_high", np.nan), errors="coerce")
    if np.isfinite(low) and np.isfinite(high):
        return f"{mean:.3f} [{float(low):.3f}, {float(high):.3f}]"
    return f"{mean:.3f}"


def selected_comparison(summary: pd.DataFrame, runs: pd.DataFrame,
                        selections: pd.DataFrame) -> pd.DataFrame:
    rows = []
    for task in ["view", "domain"]:
        external = "OAI_test" if task == "view" else "FracAtlas_test"
        chosen = selections[selections.task == task]
        identities = []
        for _, selected in chosen.iterrows():
            base_regime = selected.regime
            identities.append((selected.track, selected.model,
                               selected.domain_model, base_regime))
        c_models = summary[(summary.track == "C") & (summary.task == task) &
                           (summary.regime == "zero_shot")].model.dropna().unique()
        identities.extend(("C", m, np.nan, "zero_shot") for m in sorted(c_models))
        for track, model, domain_model, regime in identities:
            # Track A/B view exports retain the explicit UPMC prefix, whereas
            # Track C names the same internal cohort simply internal_test.
            internal = (
                "UPMC_internal_test"
                if task == "view" and track in {"A", "B"}
                else "internal_test"
            )
            for split_type, split in [("Internal", internal), ("External", external)]:
                part = summary[(summary.track == track) & (summary.task == task) &
                               (summary.model == model) & (summary.regime == regime) &
                               (summary.split == split)]
                if task == "domain" and track != "C":
                    part = part[part.domain_model.fillna("") == str(domain_model or "")]
                if len(part) != 1:
                    raise RuntimeError(
                        f"Expected one row for {track}/{task}/{model}/{regime}/{split}; "
                        f"found {len(part)}"
                    )
                r = part.iloc[0]
                out = {
                    "task": task.title(), "cohort": split_type, "split": split,
                    "track": track, "model": model, "model_display": label(model),
                    "regime": regime, "regime_display": label(regime),
                    "domain_model": domain_model,
                    "domain_model_display": label(domain_model),
                    "n_runs": int(r.n_runs),
                }
                for metric in METRICS:
                    out[metric] = r.get(f"{metric}_mean", np.nan)
                    out[f"{metric}_sd"] = r.get(f"{metric}_sd", np.nan)
                    out[f"{metric}_report"] = formatted_result(r, metric, runs)
                    run_match = runs[(runs.track == track) & (runs.task == task) &
                                     (runs.split == split) & (runs.model == model) &
                                     (runs.regime == regime)]
                    if task == "domain" and track != "C":
                        run_match = run_match[
                            run_match.domain_model.fillna("") == str(domain_model or "")
                        ]
                    if track == "C" and len(run_match) == 1:
                        out[f"{metric}_ci_low"] = run_match.iloc[0].get(
                            f"{metric}_ci_low", np.nan)
                        out[f"{metric}_ci_high"] = run_match.iloc[0].get(
                            f"{metric}_ci_high", np.nan)
                    else:
                        out[f"{metric}_ci_low"] = np.nan
                        out[f"{metric}_ci_high"] = np.nan
                rows.append(out)
    return pd.DataFrame(rows)


def all_performance(summary: pd.DataFrame, runs: pd.DataFrame) -> pd.DataFrame:
    keep_regimes = ["baseline", "full_ft", "linear_probe", "zero_shot"]
    out = summary[summary.regime.isin(keep_regimes)].copy()
    out["model_display"] = out.model.map(label)
    out["regime_display"] = out.regime.map(label)
    out["domain_model_display"] = out.domain_model.map(label)
    for metric in METRICS:
        out[f"{metric}_report"] = out.apply(
            lambda r: formatted_result(r, metric, runs), axis=1)
    return out


def maintenance_tables(summary: pd.DataFrame, deltas: pd.DataFrame,
                       selections: pd.DataFrame) -> tuple[pd.DataFrame, pd.DataFrame]:
    perf_rows = []
    delta_rows = []
    for _, selected in selections.iterrows():
        if selected.task == "view":
            part = summary[identity_mask(summary, selected) &
                           (summary.regime == "maintenance")].copy()
            if not part.empty:
                perf_rows.append(part)
        else:
            regimes = ["domain_maintenance_baseline_no_injection",
                       "domain_maintenance_fewshot_injection"]
            part = summary[identity_mask(summary, selected) &
                           summary.regime.isin(regimes)].copy()
            if not part.empty:
                perf_rows.append(part)
            d = deltas[(deltas.track == selected.track) &
                       (deltas.task == "domain") &
                       (deltas.model == selected.model) &
                       (deltas.domain_model.fillna("") == str(selected.domain_model or ""))]
            if not d.empty:
                delta_rows.append(d)
    perf = pd.concat(perf_rows, ignore_index=True) if perf_rows else pd.DataFrame()
    delta = pd.concat(delta_rows, ignore_index=True) if delta_rows else pd.DataFrame()
    for df in [perf, delta]:
        if not df.empty:
            df["model_display"] = df.model.map(label)
            if "domain_model" in df:
                df["domain_model_display"] = df.domain_model.map(label)
    return perf, delta


def latex_escape(value) -> str:
    text = str(value)
    replacements = {
        "\\": r"\textbackslash{}", "&": r"\&", "%": r"\%", "$": r"\$",
        "#": r"\#", "_": r"\_", "{": r"\{", "}": r"\}",
    }
    return "".join(replacements.get(ch, ch) for ch in text)


def write_latex(df: pd.DataFrame, path: Path, caption: str, label_name: str) -> None:
    display = df.fillna("").astype(str).map(latex_escape)
    colspec = "l" * len(display.columns)
    lines = [r"\begin{table*}[t]", r"\centering", r"\small",
             rf"\caption{{{latex_escape(caption)}}}",
             rf"\label{{{latex_escape(label_name)}}}",
             rf"\begin{{tabular}}{{{colspec}}}", r"\toprule",
             " & ".join(latex_escape(c) for c in display.columns) + r" \\",
             r"\midrule"]
    for row in display.itertuples(index=False, name=None):
        lines.append(" & ".join(row) + r" \\")
    lines += [r"\bottomrule", r"\end{tabular}", r"\end{table*}"]
    path.write_text("\n".join(lines) + "\n")


def plot_external(comparison: pd.DataFrame, metric: str, path: Path, dpi: int) -> None:
    ext = comparison[comparison.cohort == "External"].copy()
    tasks = [t for t in ["View", "Domain"] if t in set(ext.task)]
    fig, axes = plt.subplots(1, len(tasks), figsize=(7.2, 3.4), squeeze=False)
    colors = {"A": "#4472C4", "B": "#70AD47", "C": "#A5A5A5"}
    for ax, task in zip(axes[0], tasks):
        d = ext[ext.task == task].sort_values(["track", metric], ascending=[True, False])
        names = [f"{r.track}: {r.model_display}" for r in d.itertuples()]
        values = pd.to_numeric(d[metric], errors="coerce").to_numpy(float)
        sds = pd.to_numeric(d[f"{metric}_sd"], errors="coerce").to_numpy(float)
        x = np.arange(len(d))
        ax.bar(x, values, color=[colors[t] for t in d.track], width=0.72)
        ab = np.isfinite(sds)
        if ab.any():
            ax.errorbar(x[ab], values[ab], yerr=sds[ab], fmt="none", ecolor="#333333",
                        capsize=3, linewidth=1)
        lows = pd.to_numeric(d[f"{metric}_ci_low"], errors="coerce").to_numpy(float)
        highs = pd.to_numeric(d[f"{metric}_ci_high"], errors="coerce").to_numpy(float)
        ci = np.isfinite(lows) & np.isfinite(highs)
        if ci.any():
            yerr = np.vstack([values[ci] - lows[ci], highs[ci] - values[ci]])
            ax.errorbar(x[ci], values[ci], yerr=yerr, fmt="none", ecolor="#333333",
                        capsize=3, linewidth=1)
        ax.set_xticks(x, names, rotation=45, ha="right", fontsize=8)
        ax.set_ylim(0, 1.02)
        ax.set_ylabel("Supported macro-F1")
        ax.set_title(f"{task}: external performance")
        ax.grid(axis="y", color="#DDDDDD", linewidth=.6)
        ax.spines[["top", "right"]].set_visible(False)
    fig.tight_layout()
    fig.savefig(path, dpi=dpi, bbox_inches="tight")
    plt.close(fig)


def plot_shift(comparison: pd.DataFrame, metric: str, path: Path, dpi: int) -> None:
    wide = comparison.pivot_table(
        index=["task", "track", "model_display", "regime_display", "domain_model_display"],
        columns="cohort", values=metric, aggfunc="first"
    ).reset_index()
    wide = wide.dropna(subset=["Internal", "External"])
    wide["change"] = wide.External - wide.Internal
    wide["label"] = wide.track + ": " + wide.model_display
    wide = wide.sort_values(["task", "change"])
    fig, axes = plt.subplots(1, 2, figsize=(7.2, 3.5), squeeze=False)
    for ax, task in zip(axes[0], ["View", "Domain"]):
        d = wide[wide.task == task]
        ax.barh(d.label, d.change, color=np.where(d.change >= 0, "#70AD47", "#C0504D"))
        ax.axvline(0, color="#333333", linewidth=.8)
        ax.set_xlim(-1, 1)
        ax.set_title(f"{task}: external − internal")
        ax.set_xlabel("Change in supported macro-F1")
        ax.tick_params(axis="y", labelsize=8)
        ax.spines[["top", "right"]].set_visible(False)
        ax.grid(axis="x", color="#DDDDDD", linewidth=.6)
    fig.tight_layout()
    fig.savefig(path, dpi=dpi, bbox_inches="tight")
    plt.close(fig)


def plot_domain_maintenance(delta: pd.DataFrame, metric: str,
                            path: Path, dpi: int) -> None:
    col = f"{metric}_delta_mean"
    sd_col = f"{metric}_delta_sd"
    if delta.empty or col not in delta:
        return
    d = delta.copy().sort_values(["track", col])
    d["label"] = d.track + ": " + d.model_display + " + " + d.domain_model_display
    y = np.arange(len(d))
    means = pd.to_numeric(d[col], errors="coerce")
    sds = pd.to_numeric(d.get(sd_col, np.nan), errors="coerce")
    fig, ax = plt.subplots(figsize=(6.8, max(2.6, .45 * len(d) + 1.2)))
    ax.barh(y, means, color=np.where(means >= 0, "#70AD47", "#C0504D"))
    if np.isfinite(sds).any():
        ax.errorbar(means, y, xerr=sds, fmt="none", ecolor="#333333", capsize=3)
    ax.axvline(0, color="#333333", linewidth=.8)
    ax.set_yticks(y, d.label, fontsize=8)
    ax.set_xlabel("Few-shot − matched no-injection macro-F1")
    ax.set_title("External domain-maintenance effect")
    ax.grid(axis="x", color="#DDDDDD", linewidth=.6)
    ax.spines[["top", "right"]].set_visible(False)
    fig.tight_layout()
    fig.savefig(path, dpi=dpi, bbox_inches="tight")
    plt.close(fig)


def plot_view_maintenance(perf: pd.DataFrame, metric: str,
                          path: Path, dpi: int) -> None:
    if perf.empty or "shots_per_class" not in perf:
        return
    d = perf[(perf.task == "view") & (perf.split == "OAI_test") &
             perf.shots_per_class.notna()].copy()
    if d.empty:
        return
    fig, ax = plt.subplots(figsize=(6.2, 3.6))
    colors = {"A": "#4472C4", "B": "#70AD47"}
    for (track, model), g in d.groupby(["track", "model"], sort=True):
        g = g.sort_values("shots_per_class")
        x = pd.to_numeric(g.shots_per_class, errors="coerce")
        y = pd.to_numeric(g[f"{metric}_mean"], errors="coerce")
        sd = pd.to_numeric(g[f"{metric}_sd"], errors="coerce")
        ax.errorbar(x, y, yerr=sd, marker="o", capsize=3, linewidth=1.7,
                    color=colors.get(track, "#777777"),
                    label=f"Track {track}: {label(model)}")
    ax.set_ylim(0, 1.02)
    ax.set_xlabel("Maintenance examples per class")
    ax.set_ylabel("Supported macro-F1")
    ax.set_title("OAI view-classification maintenance")
    ax.legend(frameon=False, fontsize=8)
    ax.grid(color="#DDDDDD", linewidth=.6)
    ax.spines[["top", "right"]].set_visible(False)
    fig.tight_layout()
    fig.savefig(path, dpi=dpi, bbox_inches="tight")
    plt.close(fig)


def write_excel(tables: dict[str, pd.DataFrame], path: Path) -> bool:
    try:
        import openpyxl  # noqa: F401
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError:
        return False
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        for name, df in tables.items():
            df.to_excel(writer, sheet_name=name[:31], index=False)
        wb = writer.book
        for ws in wb.worksheets:
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = ws.dimensions
            ws.sheet_view.showGridLines = False
            for cell in ws[1]:
                cell.fill = PatternFill("solid", fgColor="1F4E78")
                cell.font = Font(color="FFFFFF", bold=True)
                cell.alignment = Alignment(horizontal="center", vertical="center")
            for column in ws.columns:
                letter = get_column_letter(column[0].column)
                width = min(38, max(10, max(len(str(c.value or "")) for c in column) + 2))
                ws.column_dimensions[letter].width = width
            for row in ws.iter_rows(min_row=2):
                for cell in row:
                    cell.alignment = Alignment(vertical="center")
                    if isinstance(cell.value, float):
                        cell.number_format = "0.000"
    return True


def main() -> None:
    args = parse_args()
    root = Path(args.analysis_dir).expanduser().resolve()
    out = Path(args.out_dir).expanduser().resolve()
    tables_dir, figures_dir = out / "tables", out / "figures"
    tables_dir.mkdir(parents=True, exist_ok=True)
    figures_dir.mkdir(parents=True, exist_ok=True)

    summary = read_required(root, "harmonized_model_summary.csv")
    runs = read_required(root, "harmonized_run_metrics.csv")
    deltas = read_required(root, "harmonized_maintenance_delta_summary.csv")
    check_columns(summary, ["track", "task", "split", "model", "domain_model",
                            "regime", "n_runs", f"{args.selection_metric}_mean"],
                  "harmonized_model_summary.csv")

    selections = selection_table(summary, args.selection_metric)
    comparison = selected_comparison(summary, runs, selections)
    all_results = all_performance(summary, runs)
    maintenance, selected_deltas = maintenance_tables(summary, deltas, selections)

    compact_cols = [
        "task", "cohort", "track", "model_display", "regime_display",
        "domain_model_display", "n_runs", "accuracy_report",
        "balanced_accuracy_supported_report", "macro_f1_supported_report",
        "invalid_prediction_rate_report",
    ]
    compact = comparison[compact_cols].copy()
    compact.columns = [
        "Task", "Cohort", "Track", "Model", "Regime", "Domain classifier",
        "Runs", "Accuracy", "Balanced accuracy", "Supported macro-F1",
        "Invalid prediction rate",
    ]

    exports = {
        "selection_audit": selections,
        "isbi_compact_comparison": compact,
        "selected_comparison_numeric": comparison,
        "all_baseline_results": all_results,
        "selected_maintenance": maintenance,
        "selected_domain_deltas": selected_deltas,
    }
    for name, df in exports.items():
        df.to_csv(tables_dir / f"{name}.csv", index=False)

    write_latex(
        compact, tables_dir / "isbi_compact_comparison.tex",
        "Internally selected Track A/B models and zero-shot Track C models. "
        "Track A/B values are mean (SD) across checkpoint seeds; Track C values "
        "are point estimates [cluster-bootstrap 95% CI].",
        "tab:harmonized_comparison",
    )
    workbook_written = write_excel(exports, out / "manuscript_tables.xlsx")

    plot_external(comparison, args.primary_metric,
                  figures_dir / "external_supported_macro_f1.png", args.dpi)
    plot_shift(comparison, args.primary_metric,
               figures_dir / "internal_external_change.png", args.dpi)
    plot_domain_maintenance(selected_deltas, args.primary_metric,
                            figures_dir / "domain_maintenance_delta.png", args.dpi)
    plot_view_maintenance(maintenance, args.primary_metric,
                          figures_dir / "view_maintenance_curve.png", args.dpi)

    manifest = {
        "analysis_dir": str(root),
        "selection_metric": args.selection_metric,
        "selection_used_external_results": False,
        "primary_metric": args.primary_metric,
        "track_ab_reporting": "mean (SD) across independent checkpoint seeds",
        "track_c_reporting": "single deterministic estimate with clustered bootstrap CI",
        "workbook_written": workbook_written,
        "selected_models": selections.to_dict(orient="records"),
    }
    (out / "manuscript_output_manifest.json").write_text(
        json.dumps(manifest, indent=2, default=str) + "\n"
    )
    print("Manuscript outputs created:", out)
    print("Track A/B selection used internal results only.")
    print(selections[["task", "track", "model", "regime", "domain_model",
                      "internal_selection_value"]].to_string(index=False))
    print("Workbook:", "written" if workbook_written else "not written (openpyxl unavailable)")


if __name__ == "__main__":
    main()
